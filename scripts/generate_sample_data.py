import os
import pymupdf as fitz
import docx
from PIL import Image, ImageDraw, ImageFont

SAMPLE_DIR = os.path.join(os.path.dirname(os.path.dirname(__file__)), "sample_data")
os.makedirs(SAMPLE_DIR, exist_ok=True)

def create_sample_pdf():
    pdf_path = os.path.join(SAMPLE_DIR, "qualcomm_snapdragon_whitepaper.pdf")
    doc = fitz.open()

    # Page 1
    page1 = doc.new_page(width=595, height=842)
    text1 = """QUALCOMM SNAPDRAGON® X ELITE TECHNICAL WHITEPAPER
Title: Private On-Device AI Architecture for Next-Generation Windows PCs
Author: Qualcomm Snapdragon AI Lab & Engineering Group
Date: October 2026

1. EXECUTIVE SUMMARY & MOTIVATION
Modern desktop computing is undergoing a structural paradigm shift toward on-device neural processing. Traditional cloud-based AI workflows introduce critical vulnerabilities: privacy risks for proprietary enterprise data, high latency over wireless connections, persistent monthly subscription fees, and complete failure in offline environments.

The Snapdragon® X Elite processor revolutionizes personal computing by introducing a dedicated Qualcomm® Hexagon™ NPU capable of 45 Trillion Operations Per Second (TOPS). When paired with HP Snapdragon-powered PCs, users can execute complex Retrieval-Augmented Generation (RAG), real-time vector search, and multimodal OCR locally with zero data egress.

2. CORE HARDWARE PILLARS
• Dedicated Hexagon NPU: 45 TOPS dedicated neural acceleration isolated from the CPU/GPU.
• Qualcomm Oryon™ CPU: 12 high-performance cores delivering workstation-class responsiveness with ultra-low thermal dissipation.
• Unified Memory Subsystem: Up to 64GB LPDDR5x memory operating at 8448 MT/s, eliminating PCI-e transfer bottlenecks during LLM weight caching.
• Power Efficiency: Up to 3x higher performance-per-watt compared to legacy x86 architectures, ensuring multi-day battery life during intense AI workloads.
"""
    page1.insert_text((50, 60), text1, fontsize=10, lineheight=1.4)

    # Page 2
    page2 = doc.new_page(width=595, height=842)
    text2 = """3. SOFTWARE ARCHITECTURE: QUALCOMM AI HUB & QNN EXECUTION PROVIDER
Qualcomm AI Hub serves as the developer gateway for deploying state-of-the-art foundation models directly to Snapdragon hardware.

Key software components include:
1. Qualcomm AI Runtime (QAIRT): Low-level hardware drivers interfacing directly with the Hexagon Tensor Processor (HTP).
2. ONNX Runtime QNN Execution Provider (QNN EP): Enables native execution of ONNX model graphs directly on the NPU via Windows on Snapdragon.
3. Qualcomm AI Hub Model Catalog: Pre-quantized and verified INT4/W4A16 models including Llama-3.2-1B/3B, Phi-3.5-mini, and All-MiniLM-L6-v2 embeddings.

4. PRIVACY AND LOCAL DATA SOVEREIGNTY
By confining embedding generation and token decoding to on-chip SRAM and unified LPDDR5x memory, sensitive documents—such as internal financial audits, patient health records, legal contracts, and proprietary codebases—never traverse external networks.
"""
    page2.insert_text((50, 60), text2, fontsize=10, lineheight=1.4)

    # Page 3
    page3 = doc.new_page(width=595, height=842)
    text3 = """5. BENCHMARK RESULTS & BATTERY EFFICIENCY
Empirical testing on HP Snapdragon-powered laptops demonstrates dramatic improvements over CPU-only fallbacks:
• First-Token Latency: Reduced from 420ms (x86 CPU) down to 32ms on Hexagon NPU.
• Document Embedding Throughput: 1,450 chunks per second using INT8 quantized MiniLM.
• System Power Draw: Average 4.2 Watts under continuous generation vs 38 Watts on discrete GPU laptops.
• Offline Resilience: 100% functionality maintained with zero internet connectivity.

6. CONCLUSION & STRATEGIC RECOMMENDATIONS
On-device AI workspaces represent the definitive standard for confidential knowledge management. Developers must target the QNN Execution Provider to maximize performance per watt on Snapdragon platforms.
"""
    page3.insert_text((50, 60), text3, fontsize=10, lineheight=1.4)

    # Page 4
    page4 = doc.new_page(width=595, height=842)
    text4 = """7. ACTION ITEMS & SUBMISSION DEADLINES
For the Qualcomm Snapdragon AI Lab Build & Present Challenge 2026, the team must execute the following critical milestones:

• Action Item 1: Complete end-to-end testing of local document parsing and RAG citations before October 10, 2026.
• Action Item 2: Validate QNN Execution Provider fallback mechanisms across both ARM64 Snapdragon and x86_64 host machines.
• Action Item 3: Deliver comprehensive Privacy Audit logs proving zero cloud API calls and zero network bytes transmitted.
• Action Item 4: Record final 3-minute video presentation and publish public GitHub repository by October 15, 2026.
• Action Item 5: Review benchmark metrics on target HP OmniBook device prior to final jury evaluation.
"""
    page4.insert_text((50, 60), text4, fontsize=10, lineheight=1.4)

    doc.save(pdf_path)
    doc.close()
    print(f"Created sample PDF: {pdf_path}")

def create_sample_docx():
    docx_path = os.path.join(SAMPLE_DIR, "quarterly_ai_strategy.docx")
    doc = docx.Document()
    doc.add_heading("Q4 Confidential AI Strategy & Resource Plan", level=0)

    doc.add_paragraph("Executive Summary: This document details the quarterly expansion plan for SnapAssist on HP Snapdragon AI PCs.")

    doc.add_heading("1. Strategic Milestones", level=1)
    p = doc.add_paragraph()
    p.add_run("• Milestone Alpha: ").bold = True
    p.add_run("Deploy on-device RAG engine with zero external dependencies by November 1, 2026.")

    p2 = doc.add_paragraph()
    p2.add_run("• Milestone Beta: ").bold = True
    p2.add_run("Implement local image OCR and screenshot analysis pipeline by November 15, 2026.")

    doc.add_heading("2. Assigned Action Items & Deadlines", level=1)
    table = doc.add_table(rows=1, cols=3)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Action Item"
    hdr_cells[1].text = "Lead Engineer"
    hdr_cells[2].text = "Target Date"

    items = [
        ("Finalize QNN execution provider driver integration", "Aditya (Lead AI Eng)", "November 5, 2026"),
        ("Conduct battery consumption benchmarks on HP OmniBook", "Hardware Test Team", "November 12, 2026"),
        ("Verify zero cloud data leakage audit report", "Security Officer", "November 20, 2026")
    ]

    for item, lead, dt in items:
        row_cells = table.add_row().cells
        row_cells[0].text = item
        row_cells[1].text = lead
        row_cells[2].text = dt

    doc.save(docx_path)
    print(f"Created sample DOCX: {docx_path}")

def create_sample_txt():
    txt_path = os.path.join(SAMPLE_DIR, "developer_system_log.txt")
    content = """[2026-09-30 02:15:00] [SYSTEM] Initializing SnapAssist On-Device Workspace v1.0.0
[2026-09-30 02:15:01] [HARDWARE] Detecting platform: Qualcomm Snapdragon X Series / HP PC Target
[2026-09-30 02:15:02] [NPU] Hexagon Tensor Processor detected: 45 TOPS rated compute
[2026-09-30 02:15:03] [ORT] Registering ONNX Runtime Providers: QNNExecutionProvider, CPUExecutionProvider
[2026-09-30 02:15:04] [SECURITY] Network isolation confirmed: Egress blocked, localhost mode active
[2026-09-30 02:15:05] [PRIVACY_AUDIT] Cloud API Calls: 0 | Documents Exported: 0
[2026-09-30 02:15:06] [STORAGE] SQLite local database initialized at backend/data/snapassist.db
[2026-09-30 02:15:07] [EMBEDDING] Vector dimension: 384 | High-dimensional dense cosine similarity index loaded
[2026-09-30 02:15:08] [STATUS] SnapAssist ready for private local queries.
"""
    with open(txt_path, "w", encoding="utf-8") as f:
        f.write(content)
    print(f"Created sample TXT: {txt_path}")

def create_sample_image():
    img_path = os.path.join(SAMPLE_DIR, "system_architecture_screenshot.png")
    img = Image.new("RGB", (900, 500), color=(15, 23, 42))
    draw = ImageDraw.Draw(img)

    # Draw mock UI windows / header
    draw.rectangle([(20, 20), (880, 70)], fill=(30, 41, 59), outline=(230, 0, 18), width=2)
    draw.text((40, 35), "SnapAssist — Private On-Device AI Workspace (Snapdragon NPU Active)", fill=(248, 250, 252))

    # Draw stat cards
    cards = [
        ("Local Privacy", "100% On-Device", (40, 100)),
        ("Inference Provider", "Qualcomm QNN EP", (320, 100)),
        ("Cloud API Calls", "0 (Zero Egress)", (600, 100)),
    ]
    for title, val, pos in cards:
        x, y = pos
        draw.rectangle([(x, y), (x + 250, y + 100)], fill=(30, 41, 59), outline=(71, 85, 105), width=1)
        draw.text((x + 20, y + 20), title, fill=(148, 163, 184))
        draw.text((x + 20, y + 55), val, fill=(230, 0, 18) if "0" in val or "100" in val else (255, 255, 255))

    # Draw code block / log section
    draw.rectangle([(40, 230), (860, 460)], fill=(10, 15, 29), outline=(51, 65, 85), width=1)
    draw.text((60, 250), "> ONNX Runtime loaded with QNNExecutionProvider", fill=(34, 197, 94))
    draw.text((60, 280), "> Hexagon NPU: Active (Model: Llama-3.2-1B-Instruct-QNN)", fill=(241, 245, 249))
    draw.text((60, 310), "> Query Latency: 18.4ms | Throughput: 42.6 tokens/sec", fill=(241, 245, 249))
    draw.text((60, 340), "> Target Device: HP Snapdragon X Elite Copilot+ PC", fill=(230, 0, 18))
    draw.text((60, 370), "> Privacy Audit: Verified 0 bytes transmitted outside localhost", fill=(56, 189, 248))
    draw.text((60, 400), "> Grounded RAG: Page citations attached to all responses", fill=(203, 213, 225))

    img.save(img_path)
    print(f"Created sample Image: {img_path}")

if __name__ == "__main__":
    create_sample_pdf()
    create_sample_docx()
    create_sample_txt()
    create_sample_image()
